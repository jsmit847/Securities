"""
crefc_convert.py — CREFC IRP v8.4 raw-file conversion engine.

Drop-in for app.py:
    from crefc_convert import convert, detect_type, infer_period, FILE_TYPES, CREFC_HEADERS
and for crefc_template.py:
    from crefc_convert import date_column_indices

Header resolution (no more "file not found" on deploy)
------------------------------------------------------
Headers come from the first source that works, in order:
  1. an explicit workbook at $CREFC_HEADER_FILE (if set and present);
  2. any CREFC header workbook sitting next to this module — matched by a
     glob, so both "8.4 CREFC Header File.xlsx" and "8_4_CREFC_Header_File.xlsx"
     are found;
  3. the embedded CREFC_HEADERS dict in crefc_headers.py.
That means the app runs on Streamlit Cloud even if the .xlsx isn't committed
or is named differently.

Conversion behaviour
--------------------
* Real CSV parsing (quoted address fields with commas stay in one column).
* Every date column (header contains "date", or the FINF as-of-date columns)
  is normalized YYYYMMDD -> real Excel date; placeholders become blank.
* Numeric columns are coerced to numbers; CUSIP / Zip stay text.
"""

from __future__ import annotations

import csv
import glob
import io
import os
import re
import datetime as _dt
from dataclasses import dataclass, field
from functools import lru_cache
from typing import Optional

import pandas as pd
import openpyxl

# Embedded headers committed in the repo (crefc_headers.py). Optional import so
# the module still loads if that file is ever missing.
try:
    from crefc_headers import CREFC_HEADERS as _EMBEDDED_BY_CODE
except Exception:  # noqa: BLE001
    _EMBEDDED_BY_CODE = {}


# --------------------------------------------------------------------------- #
# type-code -> display / output metadata. `sheet` matches the header-workbook
# tab name AND the on-screen tab; `out` is used for output filenames.
# --------------------------------------------------------------------------- #
FILE_TYPES: dict[str, dict] = {
    "CBND": {"label": "Bond / Certificate", "sheet": "Bond",       "out": "Bond"},
    "CCOL": {"label": "Collateral Summary", "sheet": "Coll Summ",  "out": "CollSUM"},
    "LPER": {"label": "Loan Periodic",      "sheet": "Periodic",   "out": "Periodic"},
    "PROP": {"label": "Property",           "sheet": "Property",   "out": "Property"},
    "FINF": {"label": "Financial",          "sheet": "Financial",  "out": "Financial"},
    "TLOAN": {"label": "Total Loan",        "sheet": "Total Loan", "out": "TotalLoan"},
    "CLTL":  {"label": "Total Loan",        "sheet": "Total Loan", "out": "TotalLoan"},
}
_SHEET_TO_CODE = {meta["sheet"]: code for code, meta in FILE_TYPES.items()}

_SUFFIX_ALIASES = {
    "CBND": "CBND", "BOND": "CBND",
    "CCOL": "CCOL", "COLLSUM": "CCOL", "COLL": "CCOL",
    "LPER": "LPER", "PERIODIC": "LPER",
    "PROP": "PROP", "PROPERTY": "PROP",
    "FINF": "FINF", "FIN": "FINF", "FINANCIAL": "FINF",
    "TLOAN": "TLOAN", "CLTL": "CLTL", "TOTALLOAN": "TLOAN",
}

_BLANKS = {"", " ", "--", "-", "n/a", "na", "n'a", "nan", "none", "null"}
_KEEP_TEXT = ("cusip", "zip")


# --------------------------------------------------------------------------- #
# Header resolution (workbook if found, else embedded) — all cached
# --------------------------------------------------------------------------- #
@lru_cache(maxsize=1)
def _resolve_header_path() -> Optional[str]:
    env = os.environ.get("CREFC_HEADER_FILE")
    if env and os.path.exists(env):
        return env
    here = os.path.dirname(os.path.abspath(__file__))
    for name in ("8_4_CREFC_Header_File.xlsx", "8.4 CREFC Header File.xlsx"):
        p = os.path.join(here, name)
        if os.path.exists(p):
            return p
    hits = [h for h in glob.glob(os.path.join(here, "*CREFC*Header*File*.xlsx"))
            if not os.path.basename(h).startswith("~$")]
    return hits[0] if hits else None


@lru_cache(maxsize=1)
def _headers_from_file() -> Optional[dict]:
    """sheet -> headers from the workbook, or None if no workbook is available."""
    path = _resolve_header_path()
    if not path:
        return None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: dict[str, list[str]] = {}
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        out[ws.title] = [("" if v is None else str(v)) for v in first]
    wb.close()
    return out


def _headers_for(code: str) -> list[str]:
    from_file = _headers_from_file()
    if from_file is not None:
        sheet = FILE_TYPES[code]["sheet"]
        if sheet in from_file:
            return list(from_file[sheet])
    if code in _EMBEDDED_BY_CODE:
        return list(_EMBEDDED_BY_CODE[code])
    raise KeyError(
        f"No headers for {code!r}: no workbook found (looked for a *CREFC*Header*File*.xlsx "
        f"next to crefc_convert.py or $CREFC_HEADER_FILE) and it isn't in crefc_headers.CREFC_HEADERS."
    )


class _HeaderView(dict):
    """Lazy {code: headers} so importing the module never forces a read."""
    def __missing__(self, code):
        self[code] = _headers_for(code)
        return self[code]


CREFC_HEADERS = _HeaderView()


# --------------------------------------------------------------------------- #
# Type detection & period inference
# --------------------------------------------------------------------------- #
def detect_type(name: str) -> Optional[str]:
    base = os.path.basename(name)
    stem = os.path.splitext(base)[0]
    token = re.split(r"[_\-\s]", stem)[-1].upper()
    return _SUFFIX_ALIASES.get(token)


def _yyyymmdd_to_date(s: str) -> Optional[_dt.datetime]:
    if re.fullmatch(r"\d{8}", s):
        try:
            return _dt.datetime(int(s[:4]), int(s[4:6]), int(s[6:8]))
        except ValueError:
            return None
    return None


def infer_period(df: pd.DataFrame) -> Optional[str]:
    for col in df.columns:
        if str(col).strip().lower() == "distribution date":
            for v in df[col]:
                if isinstance(v, (_dt.date, _dt.datetime)):
                    return f"{v.year:04d}_{v.month:02d}"
                d = _yyyymmdd_to_date(str(v).strip().replace(".0", ""))
                if d:
                    return f"{d.year:04d}_{d.month:02d}"
    return None


# --------------------------------------------------------------------------- #
# Cell / column helpers
# --------------------------------------------------------------------------- #
def _clean_str(v) -> str:
    return "" if v is None else str(v).strip()


def _is_date_header(h: str) -> bool:
    low = h.strip().lower()
    return ("date" in low) or (low == "yyyymmdd")


def date_column_indices(source) -> list[int]:
    """Zero-based indices of date columns. `source` may be a type code
    ('LPER'), a sheet name ('Periodic'), a header list, or a DataFrame."""
    if isinstance(source, pd.DataFrame):
        headers = [str(c) for c in source.columns]
    elif isinstance(source, (list, tuple)):
        headers = [str(c) for c in source]
    elif isinstance(source, str):
        if source in FILE_TYPES:
            headers = _headers_for(source)
        elif source in _SHEET_TO_CODE:
            headers = _headers_for(_SHEET_TO_CODE[source])
        else:
            ff = _headers_from_file() or {}
            headers = ff.get(source)
            if headers is None:
                raise KeyError(f"Unknown header source {source!r}")
    else:
        raise TypeError(f"date_column_indices: unsupported source {type(source).__name__}")
    return [i for i, h in enumerate(headers) if _is_date_header(str(h))]


def _normalize_date_series(s: pd.Series) -> pd.Series:
    def conv(v):
        t = _clean_str(v).replace(".0", "")
        if t.lower() in _BLANKS or t == "0":
            return None
        d = _yyyymmdd_to_date(t)
        return d if d is not None else (v if t else None)
    return s.map(conv)


def _coerce_numeric_series(s: pd.Series) -> Optional[pd.Series]:
    """If every non-blank value parses as a number, return an OBJECT series of
    native Python int/float/None (no numpy scalars, no pandas NA) so openpyxl
    and the template writer can consume it directly. Otherwise return None."""
    stripped = s.map(_clean_str)
    is_blank = stripped.str.lower().isin(_BLANKS)
    nonblank = stripped[~is_blank]
    if nonblank.empty:
        return None
    parsed = pd.to_numeric(nonblank.str.replace(",", "", regex=False), errors="coerce")
    if not parsed.notna().all():
        return None
    all_int = bool((parsed % 1 == 0).all())

    def conv(v):
        t = _clean_str(v)
        if t.lower() in _BLANKS:
            return None
        num = float(t.replace(",", ""))
        return int(num) if all_int else num

    return s.map(conv).astype(object)


# --------------------------------------------------------------------------- #
# Result container + main entry point
# --------------------------------------------------------------------------- #
@dataclass
class ConversionResult:
    df: pd.DataFrame
    ok: bool = True
    best_effort: bool = False
    messages: list[str] = field(default_factory=list)

    @property
    def n_rows(self) -> int:
        return len(self.df)

    @property
    def n_cols(self) -> int:
        return self.df.shape[1]


def convert(data: bytes, code: str, normalize_dates: bool = True) -> ConversionResult:
    if code not in FILE_TYPES:
        raise ValueError(f"Unknown file type code: {code!r}")

    headers = _headers_for(code)
    n = len(headers)
    messages: list[str] = []

    text = data.decode("utf-8-sig", "replace")
    rows = [r for r in csv.reader(io.StringIO(text)) if r and any(c.strip() for c in r)]

    fixed, ragged = [], 0
    for r in rows:
        if len(r) < n:
            r = r + [""] * (n - len(r))
        elif len(r) > n:
            ragged += 1
            r = r[:n]
        fixed.append(r)
    if ragged:
        messages.append(f"{ragged} row(s) had extra fields (truncated to {n} cols)")

    df = pd.DataFrame(fixed, columns=headers)
    df.columns = pd.RangeIndex(len(headers))  # address by position (dupes-safe)

    for i, h in enumerate(headers):
        if normalize_dates and _is_date_header(h):
            df[i] = _normalize_date_series(df[i])
            continue
        low = h.strip().lower()
        if any(k in low for k in _KEEP_TEXT):
            df[i] = df[i].map(lambda v: None if _clean_str(v).lower() in _BLANKS else _clean_str(v))
            continue
        num = _coerce_numeric_series(df[i])
        if num is not None:
            df[i] = num
        else:
            df[i] = df[i].map(lambda v: None if _clean_str(v).lower() in _BLANKS else _clean_str(v))

    df.columns = headers
    ok = ragged == 0
    if not messages:
        messages.append("clean")
    return ConversionResult(df=df, ok=ok, best_effort=(code == "FINF"), messages=messages)
